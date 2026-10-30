#!/usr/bin/env python3
"""Combine genome and protein BUSCO TXT summaries by accession basename."""
from pathlib import Path
import csv
import re

BASE = Path("/mnt/d/Bipolaris_genomes/Bipolaris/BUSCO")
RESULTS = BASE / "results"
RESULTS.mkdir(parents=True, exist_ok=True)
PATTERN = re.compile(
    r"C:([\d.]+)%\s*\[S:([\d.]+)%,\s*D:([\d.]+)%\],"
    r"\s*F:([\d.]+)%,\s*M:([\d.]+)%,\s*n:(\d+)"
)

def read_results(folder):
    data = {}
    if not folder.exists():
        return data
    for sample_dir in sorted(p for p in folder.iterdir() if p.is_dir()):
        summaries = list(sample_dir.glob("short_summary*.txt"))
        if not summaries:
            continue
        text = summaries[0].read_text(errors="replace")
        match = PATTERN.search(text)
        if not match:
            print("WARNING: unable to parse", summaries[0])
            continue
        c, s, d, f, m, n = match.groups()
        data[sample_dir.name] = {
            "Complete_%": float(c), "Single_copy_%": float(s),
            "Duplicated_%": float(d), "Fragmented_%": float(f),
            "Missing_%": float(m), "BUSCO_count": int(n)
        }
    return data

genome = read_results(BASE / "BUSCO_RESULTS" / "GENOME")
protein = read_results(BASE / "BUSCO_RESULTS" / "PROTEIN")

columns = [
    "Sample", "Genome_Complete_%", "Genome_Single_copy_%",
    "Genome_Duplicated_%", "Genome_Fragmented_%", "Genome_Missing_%",
    "Genome_BUSCO_count", "Protein_Complete_%", "Protein_Single_copy_%",
    "Protein_Duplicated_%", "Protein_Fragmented_%", "Protein_Missing_%",
    "Protein_BUSCO_count", "Classification"
]
rows = []
for sample in sorted(set(genome) | set(protein)):
    g, p = genome.get(sample, {}), protein.get(sample, {})
    if not g or not p:
        label = "INCOMPLETE"
    elif (g["Complete_%"] >= 95 and p["Complete_%"] >= 95
          and g["Fragmented_%"] <= 2 and p["Fragmented_%"] <= 2):
        label = "KEEP"
    elif (g["Complete_%"] < 90 or p["Complete_%"] < 90
          or g["Fragmented_%"] > 5 or p["Fragmented_%"] > 5):
        label = "REMOVE_FRAGMENTED"
    else:
        label = "REVIEW"
    row = {"Sample": sample}
    row.update({"Genome_" + k: v for k, v in g.items()})
    row.update({"Protein_" + k: v for k, v in p.items()})
    row["Classification"] = label
    rows.append(row)

tsv = RESULTS / "Bipolaris_BUSCO_Genome_Protein_Summary.tsv"
with tsv.open("w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=columns, delimiter="\t")
    writer.writeheader()
    writer.writerows(rows)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "BUSCO Summary"
ws.append(columns)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
for row in rows:
    ws.append([row.get(col, "") for col in columns])
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
for col in ws.columns:
    letter = get_column_letter(col[0].column)
    ws.column_dimensions[letter].width = min(
        max(len(str(c.value or "")) for c in col) + 2, 32
    )
xlsx = RESULTS / "Bipolaris_BUSCO_Genome_Protein_Summary.xlsx"
wb.save(xlsx)

print("Genome summaries:", len(genome))
print("Protein summaries:", len(protein))
print("Report rows:", len(rows))
print("Excel:", xlsx)
print("TSV:", tsv)
